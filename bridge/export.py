"""Export the 3Rs justification as Word or structured JSON, from the same inputs as the Markdown report."""
from __future__ import annotations

import io
import json
from datetime import date

from docx import Document

from .catalogue import AREAS, DATA_VERSION, REFERENCES, REQUIREMENTS, TIER_NAMES


def as_dict(question: str, area: str, needs: list[str], results: list[dict], choice: dict | None,
            ss: dict | None, refinements: list[str], robust: dict | None = None,
            override_reasons: dict[str, str] | None = None) -> dict:
    """Machine-readable record of a Model Finder session (no personal data)."""
    return {
        "tool": "3R Path", "catalogue_version": DATA_VERSION, "date": date.today().isoformat(),
        "question": question, "area": {"key": area, "label": AREAS[area]},
        "requirements": [{"key": n, "label": REQUIREMENTS[n]} for n in needs],
        "lowest_adequate_model": ({"key": choice["key"], "name": choice["name"], "tier": TIER_NAMES[choice["tier"]],
                                   "status": choice["status"]} if choice else None),
        "robustness": robust,
        "score_overrides": override_reasons or {},
        "ladder": [{"key": r["key"], "name": r["name"], "tier": TIER_NAMES[r["tier"]], "verdict": r["verdict"],
                    "score": r["score"], "unmet": r["gaps"], "evidence": r.get("evidence", [])} for r in results],
        "reduction": ss, "refinements": refinements, "references": REFERENCES,
        "disclaimer": "A decision aid to support discussion, not regulatory advice.",
    }


def as_json(*args, **kwargs) -> str:
    return json.dumps(as_dict(*args, **kwargs), indent=2, ensure_ascii=False)


def as_docx(question: str, area: str, needs: list[str], results: list[dict], choice: dict | None,
            ss: dict | None, refinements: list[str], robust: dict | None = None,
            override_reasons: dict[str, str] | None = None) -> bytes:
    """Word version of the justification, ready to paste into an AWERB or grant document."""
    d = Document()
    d.add_heading("3Rs justification summary", 0)
    d.add_paragraph(f"Generated with 3R Path (catalogue {DATA_VERSION}) on {date.today():%d %B %Y}. "
                    "A decision aid to support discussion, not regulatory advice.")
    d.add_heading("Research question", 1)
    d.add_paragraph(question or "(not given)")
    d.add_paragraph(f"Research area: {AREAS[area]}")
    for n in needs:
        d.add_paragraph(REQUIREMENTS[n], style="List Bullet")
    d.add_heading("Replacement", 1)
    if choice:
        d.add_paragraph(f"Lowest adequate model on the ladder: {choice['name']} ({TIER_NAMES[choice['tier']]}). "
                        f"Regulatory status: {choice['status']}")
    else:
        d.add_paragraph("No model met every requirement. Consider splitting the question so earlier steps use "
                        "non-animal methods.")
    if robust:
        d.add_paragraph(f"Robustness: {robust['label']} ({robust['stable']:.0%} of {robust['n_tests']} one-point "
                        "score changes leave the recommendation unchanged).")
    for k, why in (override_reasons or {}).items():
        d.add_paragraph(f"Score changed by user for {k}: {why or '(no reason given)'}", style="List Bullet")
    t = d.add_table(rows=1, cols=4)
    t.style = "Light Grid Accent 1"
    for c, h in zip(t.rows[0].cells, ["Tier", "Model", "Verdict", "Why"]):
        c.text = h
    for r in results:
        why = "; ".join(r["gaps"]) if r["gaps"] else (
            "Not suited to this research area" if r["base"] == 0 else "Meets the stated requirements")
        for c, v in zip(t.add_row().cells, [TIER_NAMES[r["tier"]], r["name"], r["verdict"], why]):
            c.text = v
    d.add_heading("Reduction", 1)
    d.add_paragraph((ss or {}).get("text") or "No sample size calculation entered.")
    d.add_heading("Refinement", 1)
    if refinements:
        for x in refinements:
            d.add_paragraph(x, style="List Bullet")
    else:
        d.add_paragraph("No refinements selected.")
    d.add_heading("References", 1)
    for r in REFERENCES:
        d.add_paragraph(r, style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
