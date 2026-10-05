"""Streamlit interface for the Licence Guide tab."""
from __future__ import annotations

import io
from datetime import date

import streamlit as st
from docx import Document
from docx.shared import Pt, RGBColor

from . import ui
from .catalogue import TIER_NAMES
from .examples import CREDIT, EXAMPLES, SOURCE_PAGE, SOURCE_URL
from .licence import (CATEGORY_TEXT, CHECKLIST, CHECKLIST_PIL, ENDPOINTS, EXAMPLES_NOT_REGULATED, EXAMPLES_REGULATED,
                      GLOSSARY, JOURNEY, MODULES, NAMED_PERSONS, NTS, SEVERITY, SPECIES, STEP_STAGE, SUBJECTS,
                      TECHNIQUES, WORRIES, licence_verdict, modules_for, pil_categories, readability)

STEPS = ["1. Do I need a licence?", "2. Which licence and training?", "3. Who can help me?",
         "4. Prepare my application", "5. Ready to submit?"]
SAY = [
    "Hi, I'm Dot. Feeling unsure is normal: everyone starts here. Tell me what your work involves, and I'll tell you whether UK law covers it.",
    "Tick the techniques you will do with your own hands. I'll work out your personal licence category and the training you need.",
    "You are not on your own. Every licensed establishment has these people, and their job is to help you. Talk to them early.",
    "Most new researchers apply for a personal licence and work under their PI's project licence. Pick your route and I'll help you prepare.",
    "Nearly there. Tick off each item as you finish it, and see what happens after you submit.",
]


def _prefill(results, choice, ss, refinements):
    """Turn Model Finder results into first-draft text for the 3Rs sections."""
    considered = [r for r in results if r["tier"] < 5]
    alts = "; ".join(r["name"] for r in considered) or "None recorded yet."
    use_first = [r["name"] for r in considered if r["verdict"] == "adequate"]
    reasons = []
    for r in considered:
        if r["verdict"] == "adequate":
            continue
        if r["base"] == 0:
            why = "not suited to this research area"
        elif r["gaps"]:
            why = "cannot provide: " + "; ".join(g[0].lower() + g[1:] for g in r["gaps"])
        else:
            why = "only partly answers questions in this area"
        reasons.append(f"{r['name']} ({why})")
    not_suitable = ""
    if use_first:
        not_suitable += ("We will use " + ", ".join(use_first) + " first, for the parts of the question they can "
                         "answer, and use animals only for what they cannot. ")
    if reasons:
        not_suitable += "Other alternatives could not answer the question: " + "; ".join(reasons) + "."
    text = {
        "alternatives": f"We considered {alts}.",
        "not_suitable": not_suitable.strip(),
        "guidance": "PREPARE guidelines for planning, ARRIVE 2.0 for reporting, NC3Rs resources, and LASA guidance on good practice.",
        "informed": "Through the establishment's Named Information Officer, NC3Rs newsletters and events, and local 3Rs seminars.",
    }
    if ss:
        text["numbers"] = (f"We used a power calculation: to detect a difference of {ss['diff']} with a standard deviation "
                           f"of {ss['sd']}, at a significance level of {ss['alpha']} and {ss['power']:.0%} power, "
                           f"we need {ss['n']} animals per group.")
    if refinements:
        text["refine"] = "We will: " + "; ".join(r[0].lower() + r[1:] for r in refinements) + "."
    return text


def _docx(fields: dict, title: str, protocol: dict) -> bytes:
    doc = Document()
    st_ = doc.styles["Normal"]; st_.font.name = "Calibri"; st_.font.size = Pt(11)
    for h in ("Heading 1", "Heading 2"):
        doc.styles[h].font.color.rgb = RGBColor(0, 0, 0)
    doc.add_heading(title or "Project licence: draft non-technical summary", level=1)
    doc.add_paragraph().add_run(f"Draft prepared with 3R Bridge on {date.today():%d %B %Y}. "
                                "For discussion with the NTCO, NVS and AWERB before submission on ASPeL.").italic = True
    for section, items in NTS:
        doc.add_heading(section, level=2)
        for key, q, _ in items:
            p = doc.add_paragraph(); p.add_run(q).bold = True
            doc.add_paragraph(fields.get(key, "").strip() or "[to complete]")
    if protocol.get("name"):
        doc.add_heading("Protocol outline", level=2)
        for label, val in protocol.items():
            p = doc.add_paragraph(); p.add_run(label.capitalize() + ": ").bold = True
            p.add_run(val if isinstance(val, str) else "; ".join(val) or "[to complete]")
    buf = io.BytesIO(); doc.save(buf); return buf.getvalue()


ROUTES = ["My personal licence (PIL)", "A project licence summary (PPL)", "Real examples to learn from"]


def _go(i: int):
    st.session_state["guide_step"] = STEPS[i]


def _nav(i: int):
    """Back and next buttons under each step."""
    st.write("")
    b1, _, b2 = st.columns([1, 2, 1.6])
    if i > 0:
        b1.button("Back", key=f"back_{i}", on_click=_go, args=(i - 1,), width="stretch")
    if i < len(STEPS) - 1:
        b2.button("Next step", key=f"next_{i}", on_click=_go, args=(i + 1,), type="primary", width="stretch")


def _jargon():
    with st.popover("Jargon buster", icon=":material/menu_book:"):
        q = st.text_input("Look up a term", placeholder="For example: NACWO, Schedule 1, humane endpoint", key="g_term")
        hits = {k: v for k, v in GLOSSARY.items() if not q or q.lower() in (k + " " + v).lower()}
        if not hits:
            st.caption("No match. Try a shorter word, or ask your NTCO.")
        for k, v in hits.items():
            st.html(f'<p class="rb-term"><b>{k}</b>: {v}</p>')


def _picked_techniques():
    return [k for k in TECHNIQUES if st.session_state.get(f"g_tech_{k}")]


def _pil_docx(info: dict, techniques: list[str], cats: list[str], modules: list[dict], questions: str) -> bytes:
    doc = Document()
    st_ = doc.styles["Normal"]; st_.font.name = "Calibri"; st_.font.size = Pt(11)
    for h in ("Heading 1", "Heading 2"):
        doc.styles[h].font.color.rgb = RGBColor(0, 0, 0)
    doc.add_heading("Personal licence preparation sheet", level=1)
    doc.add_paragraph().add_run(f"Prepared with 3R Bridge on {date.today():%d %B %Y}. Bring this to your first meeting "
                                "with the NTCO; the application itself is made on ASPeL.").italic = True
    doc.add_heading("About me", level=2)
    for label, val in info.items():
        p = doc.add_paragraph(); p.add_run(label + ": ").bold = True
        p.add_run((", ".join(val) if isinstance(val, list) else val) or "[to complete]")
    doc.add_heading("Techniques I expect to perform", level=2)
    for t in techniques or ["[to complete]"]:
        doc.add_paragraph(t, style="List Bullet")
    p = doc.add_paragraph(); p.add_run("Likely personal licence categories: ").bold = True
    p.add_run(", ".join(cats) or "[ask the NTCO]")
    doc.add_heading("Training modules", level=2)
    table = doc.add_table(rows=1, cols=4); table.style = "Table Grid"
    for c, h in zip(table.rows[0].cells, ["Module", "Status", "Provider", "Date"]):
        c.text = h; c.paragraphs[0].runs[0].bold = True
    for m in modules:
        row = table.add_row().cells
        for c, k in zip(row, ["Module", "Status", "Provider", "Date"]):
            v = m.get(k)
            c.text = "" if v is None else (f"{v:%d %b %Y}" if hasattr(v, "strftime") else str(v))
    doc.add_heading("Questions for the NTCO", level=2)
    doc.add_paragraph(questions.strip() or "[to complete]")
    buf = io.BytesIO(); doc.save(buf); return buf.getvalue()


KEEP = ("g_", "pil_", "nts_", "prot_", "chk_")


def render_guide(results, choice, ss, refinements, question):
    # Streamlit forgets widgets that are not on screen; re-saving them keeps answers when moving between steps.
    for k in list(st.session_state):
        if isinstance(k, str) and k.startswith(KEEP) and not k.startswith("pil_mods_"):
            st.session_state[k] = st.session_state[k]
    st.session_state.setdefault("guide_step", STEPS[0])
    top1, top2 = st.columns([4, 1], vertical_alignment="center")
    top1.markdown("A plain-English guide for researchers new to in vivo work in the UK. "
                  "Work through the steps in order, or jump to the one you need.")
    with top2:
        _jargon()
    step = st.segmented_control("Step", STEPS, label_visibility="collapsed", key="guide_step") or STEPS[0]
    i = STEPS.index(step)
    ui.journey(JOURNEY, STEP_STAGE[i])
    ui.say(SAY[i])

    # ---------------------------------------------------------- step 1
    if i == 0:
        with st.expander("Not sure what counts? See everyday examples", expanded=True):
            ui.examples(EXAMPLES_REGULATED, EXAMPLES_NOT_REGULATED)
            st.caption("The legal test: could it cause pain, suffering, distress or lasting harm equal to a needle "
                       "insertion or more? If yes, it is a regulated procedure.")
        subject = st.radio("What will your work involve?", list(SUBJECTS), format_func=SUBJECTS.get, key="g_subject")
        stage, above, ga, s1 = True, False, False, False
        if subject in ("fish_amph", "mammal_bird_reptile", "cephalopod"):
            if subject == "fish_amph":
                stage = st.radio("Life stage", ["Before they can feed independently", "Able to feed independently or older"],
                                 index=1, key="g_stage_f") != "Before they can feed independently"
            elif subject == "mammal_bird_reptile":
                stage = st.radio("Life stage", ["Embryo or fetus before the last third of gestation or incubation",
                                                "Later fetal stage, born or hatched animals"], index=1, key="g_stage_m").startswith("Later")
            st.markdown("**What will you do?**")
            above = st.checkbox("Something that could cause pain, suffering, distress or lasting harm equal to a "
                                "needle insertion or more (injections, blood sampling, surgery, special diets, stressful tests)",
                                key="g_above")
            ga = st.checkbox("Breed genetically altered animals that may suffer harm from the alteration", key="g_ga")
            s1 = st.checkbox("Only kill animals humanely by a Schedule 1 method to collect tissue", key="g_s1")
        headline, level, notes = licence_verdict(subject, stage, above, ga, s1)
        ui.verdict(headline, level, notes)
        with st.expander("Where this comes from"):
            st.markdown("A procedure is regulated if it is carried out on a protected animal for a scientific or "
                        "educational purpose and may cause pain, suffering, distress or lasting harm equivalent to, or "
                        "higher than, inserting a hypodermic needle according to good veterinary practice. Protected "
                        "animals are living vertebrates other than humans, and living cephalopods. Source: Home Office, "
                        "Guidance on the operation of the Animals (Scientific Procedures) Act 1986.")

    # ---------------------------------------------------------- step 2
    elif i == 1:
        st.caption("Tick every technique you expect to perform yourself. Not sure? Ask your supervisor what the "
                   "protocols you will work on involve.")
        picked = [k for k, (label, _) in TECHNIQUES.items() if st.checkbox(label, key=f"g_tech_{k}")]
        cats = pil_categories(picked)
        mods = modules_for(cats)
        c1, c2 = st.columns(2, gap="large")
        with c1:
            ui.licence_card(cats, CATEGORY_TEXT)
        with c2:
            ui.modules_card(mods, MODULES)
            if "D" in cats:
                st.caption("Category D has additional requirements: ask your NTCO.")
        st.info("After the modules, you work under supervision until your establishment assesses you as competent "
                "for each technique and species. Your competence is reviewed regularly, at least every five years. "
                "Previous training elsewhere may count: ask your NTCO about exemptions.")

    # ---------------------------------------------------------- step 3
    elif i == 2:
        ui.people(NAMED_PERSONS)
        st.markdown("#### Common worries")
        st.caption("Questions new researchers often feel shy asking.")
        for q, a in WORRIES:
            with st.expander(q):
                st.markdown(a)

    # ---------------------------------------------------------- step 4
    elif i == 3:
        route = st.radio("What are you preparing?", ROUTES, horizontal=True, key="g_route",
                         captions=["Most new researchers start here", "Usually written by the PI", "Three published summaries"])
        if route == ROUTES[0]:
            _pil_form()
        elif route == ROUTES[2]:
            _examples()
        else:
            _ppl_form(results, choice, ss, refinements, question)

    # ---------------------------------------------------------- step 5
    else:
        pil = st.session_state.get("g_route", ROUTES[0]) != ROUTES[1]
        items = CHECKLIST_PIL if pil else CHECKLIST
        st.caption("Checklist for " + ("your personal licence." if pil else "a project licence application.")
                   + " Change the route in step 4.")
        done = [c for c in items if st.checkbox(c, key=f"chk_{c}")]
        st.progress(len(done) / len(items), text=f"{len(done)} of {len(items)} done")
        if len(done) == len(items):
            ui.say("Everything is in place. Well done, and good luck with your application.")
            if not st.session_state.get("celebrated"):
                st.balloons()
                st.session_state["celebrated"] = True
        else:
            st.session_state["celebrated"] = False
        st.markdown("#### After you submit")
        ui.journey(JOURNEY, 4)
        st.caption("Once your licence is granted, you practise under supervision until you are signed off for each "
                   "technique. Keep a training record: your NTCO will review it with you.")

    _nav(i)
    st.html('<p class="rb-foot">A learning aid based on Home Office guidance. It is not legal advice: your '
            "establishment's Named Persons and the Home Office always have the final word.</p>")


def _pil_form():
    st.caption("A preparation sheet to take to your NTCO. It gathers what the personal licence application on "
               "ASPeL will ask about, so the meeting is quick and you know what is missing.")
    c1, c2 = st.columns(2, gap="large")
    with c1:
        name = st.text_input("Your name", key="pil_name")
        role = st.text_input("Your role", key="pil_role", placeholder="For example: PhD student, postdoctoral researcher")
        sup = st.text_input("Supervisor and project licence holder", key="pil_sup")
        ppl = st.text_input("Project licence number, if you know it", key="pil_ppl")
    with c2:
        species = st.multiselect("Species", SPECIES, key="pil_species")
        exp = st.text_area("Previous experience with animals", key="pil_exp", height=122,
                           placeholder="Training or licences held elsewhere, techniques practised, species")
    picked = _picked_techniques()
    cats = pil_categories(picked)
    tech_labels = [TECHNIQUES[k][0] for k in picked]
    if tech_labels:
        st.markdown("**Techniques from step 2:** " + "; ".join(tech_labels))
    else:
        st.info("Tick your techniques in step 2 and they will appear here, with your categories and modules.")
    st.markdown("**Training modules**")
    rows = [{"Module": f"{m}: {MODULES[m]}", "Status": "Not started", "Provider": "", "Date": None}
            for m in modules_for(cats)] or [{"Module": "", "Status": "Not started", "Provider": "", "Date": None}]
    mods = st.data_editor(rows, key=f"pil_mods_{'-'.join(cats)}", num_rows="dynamic", width="stretch", hide_index=True,
                          column_config={
                              "Status": st.column_config.SelectboxColumn(options=["Not started", "Booked", "Completed", "Exemption requested"]),
                              "Date": st.column_config.DateColumn(format="DD MMM YYYY")})
    qs = st.text_area("Questions for the NTCO", key="pil_qs", height=100,
                      placeholder="For example: Does my training abroad count? Who will supervise my first procedures?")
    info = {"Name": name, "Role": role, "Supervisor and project licence holder": sup,
            "Project licence number": ppl, "Species": species, "Previous experience": exp}
    st.caption("Tip: click outside the last box you edited before downloading, so it is included.")
    st.download_button("Download preparation sheet (Word)", _pil_docx(info, tech_labels, cats, mods, qs),
                       file_name="personal_licence_preparation.docx",
                       mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


def _ppl_form(results, choice, ss, refinements, question):
    st.caption("Builds a first draft of the non-technical summary, using the headings the Home Office publishes, "
               "plus a protocol outline. Take the draft to your NTCO, NVS and AWERB.")
    if st.button("Fill the 3Rs sections from my Model Finder results"):
        for k, v in _prefill(results, choice, ss, refinements).items():
            st.session_state[f"nts_{k}"] = v
    if choice:
        st.caption(f"Model Finder suggested starting with: {choice['name']} ({TIER_NAMES[choice['tier']]}).")
    title = st.text_input("Project title", value=question, key="nts_title")

    fields = {}
    for section, items in NTS:
        with st.expander(section, expanded=section == "Objectives and benefits"):
            for key, q, hint in items:
                fields[key] = st.text_area(q, key=f"nts_{key}", placeholder=hint, height=90)

    st.markdown("#### Protocol outline")
    p1, p2 = st.columns(2, gap="large")
    with p1:
        pname = st.text_input("Protocol name", key="prot_name", placeholder="For example: Optic nerve crush and intravitreal injection")
        steps_txt = st.text_area("Steps, in order", key="prot_steps", height=120,
                                 placeholder="1. Anaesthesia ...\n2. Procedure ...\n3. Recovery and monitoring ...")
        sev = st.selectbox("Severity limit", list(SEVERITY), index=3, key="prot_sev")
        st.caption(SEVERITY[sev])
    with p2:
        ends = st.multiselect("Humane endpoints", ENDPOINTS, key="prot_ends")
        mon = st.text_input("Monitoring frequency", key="prot_mon", placeholder="For example: twice daily for 3 days after surgery, then daily")
    protocol = {"name": pname, "steps": steps_txt, "severity limit": sev, "humane endpoints": ends, "monitoring": mon}

    st.markdown("#### Plain-language check")
    lay = " ".join(fields.get(k, "") for k in ("aim", "why", "benefit", "typical"))
    if lay.strip():
        r = readability(lay)
        m1, m2, m3 = st.columns(3)
        m1.metric("Words", r["words"])
        m2.metric("Average sentence length", f"{r['avg']:.0f} words", help="Aim for under 20.")
        m3.metric("Jargon found", len(r["jargon"]))
        for w, alt in r["jargon"]:
            st.markdown(f"- Consider replacing **{w}** with *{alt}*")
    else:
        st.caption("Write the aim, importance, benefits and typical procedures above, and this check will run on them.")

    st.caption("Tip: click outside the last box you edited before downloading, so it is included.")
    st.download_button("Download draft (Word)", _docx(fields, title, protocol),
                       file_name="draft_non_technical_summary.docx",
                       mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


def _examples():
    st.caption("Every project licence has a published plain-English summary. These three are typical of the "
               "projects new researchers join. Read one alongside your own draft.")
    labels = [e["label"] for e in EXAMPLES]
    pick = st.segmented_control("Example", labels, default=labels[0], key="g_example",
                                label_visibility="collapsed") or labels[0]
    ex = EXAMPLES[labels.index(pick)]
    ui.example_card(ex)
    for i, (heading, text, tip) in enumerate(ex["sections"]):
        with st.expander(heading, expanded=i == 0):
            st.markdown(text)
            ui.lesson(tip)
    st.caption(f"{CREDIT} Read the [full summaries]({SOURCE_URL}) or [all 2025 summaries]({SOURCE_PAGE}).")
